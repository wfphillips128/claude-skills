"""Shadow model template — a Python mirror of a 5G LAMBDA, for tie-out.

Copy this file, then edit the four blocks marked  ### EDIT ###.

The point is to PROVE the LAMBDA's maths before anyone pastes it into Excel. So this
must be a *translation* of the formula, step for step: same variable names, same order,
same rounding, same signs. If you write better maths here, a passing tie-out proves
nothing about the formula you're shipping.

Read-only: the workbook is never modified.

Usage:
    python shadow_model.py <workbook.xlsx> [--benchmark <file.csv>]
"""

import argparse
import math
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  python -m pip install openpyxl")


# --------------------------------------------------------------------------------
# Tolerances. State these in the output; don't widen them to make a run pass.
# --------------------------------------------------------------------------------
DOLLAR_TOL = 0.01
RATE_TOL = 1e-9


# ================================================================================
# ### EDIT 1 ### Read the real inputs out of the workbook.
# ================================================================================
def load_inputs(path):
    """Pull the per-period input arrays from the model.

    data_only=True gives the values Excel last calculated. It returns None for every
    formula cell if the file has never been opened and saved by Excel — if that
    happens, say so rather than substituting zeros.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Bond Model"]                       # <- sheet name

    first_row, last_row = 18, 35                # <- the period rows
    col = {"cfads": 10, "par": 15, "interest": 16, "debt_service": 19}   # J, O, P, S

    def column(c):
        return [ws.cell(row=r, column=c).value for r in range(first_row, last_row + 1)]

    data = {k: column(c) for k, c in col.items()}

    missing = [k for k, v in data.items() if any(x is None for x in v)]
    if missing:
        raise SystemExit(
            f"No cached values for: {', '.join(missing)}.\n"
            "Open the workbook in Excel, save it, and re-run — openpyxl cannot "
            "evaluate formulas itself."
        )

    # Scalars from defined names or fixed cells
    data["dscr"] = ws["D14"].value
    data["par_incr"] = ws["D42"].value           # PV_Increment
    return data


# ================================================================================
# ### EDIT 2 ### Mirror the inner convergence recursion.
# ================================================================================
def solve_interest(cfads, dscr, principal, rate, par_incr,
                   tol=DOLLAR_TOL, max_iter=100):
    """Mirror of the convergence LAMBDA (e.g. fx.InterestDebtIssuanceλ).

        Payment,   IF(DoNotUse = 0, CFADS / DSCR, DoNotUse)
        Interest,  Principal * InterestRate
        Repayment, -FLOOR(CFADS / DSCR - Interest, ParValueIncr)
        Result,    IF(ABS(Payment + Repayment) < 0.01, MAX(Interest, 0), recurse(-Repayment))

    Written as a loop, not recursion, so the iteration count is inspectable.
    """
    payment = cfads / dscr                       # DoNotUse = 0 on the first call
    for i in range(1, max_iter + 1):
        interest = principal * rate
        repayment = -math.floor((cfads / dscr - interest) / par_incr) * par_incr
        if abs(payment + repayment) < tol:
            return max(interest, 0.0), i
        payment = -repayment
    # Not a bug in this script — it means the LAMBDA returns #NUM! for these inputs.
    raise RuntimeError(
        f"interest did not converge in {max_iter} iterations "
        f"(principal={principal:,.2f}, rate={rate:.6f}, increment={par_incr:,.0f}). "
        "The LAMBDA would return #NUM!. Check tolerance vs. rounding increment."
    )


# ================================================================================
# ### EDIT 3 ### Mirror the outer period recursion.
# ================================================================================
def build_schedule(inputs, opening_balance=0.0):
    """Mirror of the recursive-HSTACK LAMBDA. One dict per period.

    Excel is 1-based; `period` is kept 1-based and indexed with `period - 1` so the
    code lines up with INDEX(Array, 1, Period).
    """
    cfads = inputs["cfads"]
    dscr = inputs["dscr"]
    par_incr = inputs["par_incr"]
    n = len(cfads)

    rows = []
    principal = opening_balance
    for period in range(1, n + 1):
        i = period - 1

        rate = inputs["rate"][i] if "rate" in inputs else 0.0
        interest, iters = solve_interest(cfads[i], dscr, principal, rate, par_incr)
        payment = math.floor((cfads[i] / dscr - interest) / par_incr) * par_incr
        closing = principal - payment

        rows.append({
            "period": period,
            "opening": principal,
            "interest": interest,
            "payment": payment,
            "closing": closing,
            "debt_service": payment + interest,
            "iterations": iters,
        })
        principal = closing          # carry state — mirrors the recursive first argument

    return rows


# ================================================================================
# ### EDIT 4 ### Load the benchmark to reconcile against.
# ================================================================================
def load_benchmark(path):
    """Benchmark schedule: period -> value.

    From CSV, or parsed out of a PDF with pypdf. Whatever the source, PRINT WHAT YOU
    PARSED before comparing — a mis-parsed benchmark gives a confident wrong tie-out.
    Check the row count against the expected period count and the parsed total against
    the printed total.
    """
    import csv

    def to_num(raw):
        """Separator-safe: the rightmost of "." and "," is the decimal point.

        Stripping commas outright would read a European "1.234,56" as 1.23456 -
        a 1000x error in the very figure the tie-out trusts.
        """
        s = str(raw).strip()
        for junk in ("$", "€", "£", " ", " ", "'"):
            s = s.replace(junk, "")
        neg = s.startswith("(") and s.endswith(")")
        if neg:
            s = s[1:-1]
        last_c, last_d = s.rfind(","), s.rfind(".")
        if last_c >= 0 and last_d >= 0:
            dec = "," if last_c > last_d else "."
        elif last_c >= 0:
            dec = "." if (s.count(",") > 1 or len(s) - last_c - 1 == 3) else ","
        elif last_d >= 0:
            dec = "," if s.count(".") > 1 else "."
        else:
            dec = "."
        s = s.replace("," if dec == "." else ".", "")
        if dec != ".":
            s = s.replace(dec, ".")
        f = float(s)
        return -f if neg else f

    # utf-8-sig to match every other CSV reader in these skills: Excel's
    # "CSV UTF-8" export writes a BOM, which would break the first header key.
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return {int(r["period"]): to_num(r["value"])
                for r in csv.DictReader(fh)}


# --------------------------------------------------------------------------------
# Tie-out. Do not edit — this is the part that has to stay honest.
# --------------------------------------------------------------------------------
def tie_out(rows, benchmark, field, label, tol=DOLLAR_TOL):
    """Print a full period-by-period variance table and return True if it ties."""
    print(f"\nTie-out: shadow model vs benchmark — {label}   tolerance ${tol:,.2f}\n")
    print(f"{'Period':>6}  {'Model':>18}  {'Benchmark':>18}  {'Diff':>14}")
    print("-" * 62)

    fails, model_total, bench_total = [], 0.0, 0.0
    for r in rows:
        p = r["period"]
        model = r[field]
        bench = benchmark.get(p)
        model_total += model
        if bench is None:
            print(f"{p:>6}  {model:>18,.2f}  {'(missing)':>18}  {'':>14}")
            fails.append(p)
            continue
        bench_total += bench
        diff = model - bench
        flag = "" if abs(diff) <= tol else "  <-- FAIL"
        if flag:
            fails.append(p)
        print(f"{p:>6}  {model:>18,.2f}  {bench:>18,.2f}  {diff:>14,.2f}{flag}")

    print("-" * 62)
    total_diff = model_total - bench_total
    tflag = "" if abs(total_diff) <= tol else "  <-- FAIL"
    print(f"{'Total':>6}  {model_total:>18,.2f}  {bench_total:>18,.2f}  "
          f"{total_diff:>14,.2f}{tflag}")

    extra = sorted(set(benchmark) - {r["period"] for r in rows})
    if extra:
        print(f"\n!! Benchmark has {len(extra)} period(s) the model does not: {extra}")
        fails.extend(extra)

    ok = not fails and abs(total_diff) <= tol
    print(f"\n{'PASS' if ok else 'FAIL'} — {len(rows) - len([f for f in fails if f])}"
          f" of {len(rows)} periods within tolerance")
    if not ok:
        print("Do not deliver the LAMBDA. Fix the formula and re-run.")
    return ok


def print_schedule(rows):
    print(f"\n{'Period':>6}  {'Opening':>16}  {'Interest':>14}  {'Payment':>16}  "
          f"{'Closing':>16}  {'Iter':>4}")
    print("-" * 82)
    for r in rows:
        print(f"{r['period']:>6}  {r['opening']:>16,.2f}  {r['interest']:>14,.2f}  "
              f"{r['payment']:>16,.2f}  {r['closing']:>16,.2f}  {r['iterations']:>4}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--benchmark", help="CSV with columns: period,value")
    ap.add_argument("--field", default="payment", help="schedule field to reconcile")
    ap.add_argument("--label", default="base case")
    args = ap.parse_args()

    inputs = load_inputs(args.workbook)
    rows = build_schedule(inputs)
    print_schedule(rows)

    if args.benchmark:
        ok = tie_out(rows, load_benchmark(args.benchmark), args.field, args.label)
        sys.exit(0 if ok else 1)
    else:
        print("\nNo benchmark supplied — schedule shown but NOT verified. "
              "The LAMBDA is not ready to deliver.")


if __name__ == "__main__":
    main()
